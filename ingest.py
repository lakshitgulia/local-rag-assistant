"""
Incremental ingestion: scan a local folder, parse PDF/PPTX/DOCX/EML/MBOX/
scanned-image files, chunk the text, embed it via Ollama, and persist a FAISS
index + chunk metadata under ./index/. (No separate BM25 index file:
retrieval.py rebuilds BM25Retriever fresh from chunks.json at app startup --
that's pure CPU, no network call, and cheap enough not to need its own
persistence format.)

A content-hash manifest (index/manifest.json) means unchanged files are
skipped on every run -- only new or changed files are parsed/chunked/embedded.
A changed file has its old chunks removed from the FAISS index before the new
ones are added, so re-ingesting never leaves stale entries behind. This is
what makes watcher.py's per-file, run-on-every-fs-event calls cheap.

Run on demand:
    python ingest.py                      # scans LOCAL_SCAN_DIR from .env
    python ingest.py --scan-dir ~/Desktop/demo-docs
"""
import argparse
import hashlib
import mailbox
import os
import uuid
from email import message_from_bytes
from pathlib import Path
from typing import List, Tuple

import pytesseract
from dotenv import load_dotenv
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from pdf2image import convert_from_path
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from docx import Document as DocxDocument

from core import FAISS_DIR, assign_role, get_embeddings, load_chunks, load_manifest, save_chunks, save_manifest

load_dotenv()

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}
# Character-based (not word-based) so a PDF with broken text extraction --
# glyphs run together with no spaces -- can't produce one giant "word" that
# blows past Ollama's embedding context window.
CHUNK_CHARS = 2000  # ~400-500 tokens of normal English, ~10-15% overlap band
OVERLAP_RATIO = 0.13
MIN_OCR_CHARS = 20  # below this, a PDF page or image is treated as non-text and flagged


def chunk_text(text: str) -> List[str]:
    text = text.strip()
    if not text:
        return []
    step = int(CHUNK_CHARS * (1 - OVERLAP_RATIO))
    chunks = []
    for start in range(0, len(text), step):
        chunk = text[start:start + CHUNK_CHARS].strip()
        if chunk:
            chunks.append(chunk)
        if start + CHUNK_CHARS >= len(text):
            break
    return chunks


def parse_pdf(path: Path) -> List[Tuple[str, str]]:
    """Returns [(text, 'page N'), ...]. Falls back to OCR for pages with
    little/no extractable text (scanned pages)."""
    results = []
    reader = PdfReader(str(path))
    ocr_pages = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if len(text) < MIN_OCR_CHARS:
            ocr_pages.append(i)
        else:
            results.append((text, f"page {i + 1}"))

    if ocr_pages:
        try:
            images = convert_from_path(str(path))
            for i in ocr_pages:
                if i < len(images):
                    ocr_text = pytesseract.image_to_string(images[i]).strip()
                    if len(ocr_text) >= MIN_OCR_CHARS:
                        results.append((ocr_text, f"page {i + 1} (OCR)"))
                    else:
                        print(f"[ingest] flagged, no extractable text (visual/photo page?): {path.name} page {i + 1} -- skipped")
        except Exception as e:
            print(f"[ingest] OCR fallback failed for {path.name}: {e}")
    return results


def parse_pptx(path: Path) -> List[Tuple[str, str]]:
    results = []
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides):
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        if texts:
            results.append(("\n".join(texts), f"slide {i + 1}"))
    return results


def parse_docx(path: Path) -> List[Tuple[str, str]]:
    doc = DocxDocument(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [(text, "document")] if text.strip() else []


def parse_eml(path: Path) -> List[Tuple[str, str]]:
    with open(path, "rb") as f:
        msg = message_from_bytes(f.read())
    return _parse_email_message(msg)


def parse_mbox(path: Path) -> List[Tuple[str, str]]:
    results = []
    box = mailbox.mbox(str(path))
    for i, msg in enumerate(box):
        for text, _ in _parse_email_message(msg):
            results.append((text, f"message {i + 1}"))
    return results


def _parse_email_message(msg) -> List[Tuple[str, str]]:
    subject = msg.get("Subject", "(no subject)")
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain" and not part.get_filename():
                payload = part.get_payload(decode=True)
                if payload:
                    body += payload.decode(errors="ignore")
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            body = payload.decode(errors="ignore")
    text = f"Subject: {subject}\n\n{body}".strip()
    return [(text, "email")] if body.strip() else []


def parse_image(path: Path) -> List[Tuple[str, str]]:
    text = pytesseract.image_to_string(Image.open(path)).strip()
    if len(text) < MIN_OCR_CHARS:
        print(f"[ingest] flagged, no extractable text (visual/photo content?): {path.name} -- skipped, out of scope for this demo")
        return []
    return [(text, "image (OCR)")]


PARSERS = {
    ".pdf": parse_pdf,
    ".pptx": parse_pptx,
    ".docx": parse_docx,
    ".eml": parse_eml,
    ".mbox": parse_mbox,
}


def collect_files(scan_dir: Path) -> List[Path]:
    exts = set(PARSERS.keys()) | IMAGE_EXTS
    return [p for p in scan_dir.rglob("*") if p.is_file() and p.suffix.lower() in exts]


def file_hash(path: Path) -> str:
    """SHA-256 of file content -- the manifest key for skip-if-unchanged."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


def build_chunks(files: List[Path]) -> List[Document]:
    all_chunks = []
    for path in files:
        ext = path.suffix.lower()
        parser = PARSERS.get(ext, parse_image if ext in IMAGE_EXTS else None)
        if parser is None:
            continue
        try:
            sections = parser(path)
        except Exception as e:
            print(f"[ingest] failed to parse {path.name}: {e}")
            continue

        role = assign_role(str(path))
        for text, location in sections:
            for piece in chunk_text(text):
                all_chunks.append(Document(
                    page_content=piece,
                    # source_path (full path) identifies exactly which file a
                    # chunk came from for incremental re-indexing; source_file
                    # (basename) is what citations show -- two different
                    # files can share a basename, so these must stay distinct.
                    metadata={"source_file": path.name, "source_path": str(path), "location": location, "role": role},
                ))
        print(f"[ingest] parsed {path.name} ({len(sections)} section(s))")
    return all_chunks


EMBED_BATCH_SIZE = 64  # LangChain's embed_documents() sends the whole list in
# one HTTP request; handing it all 882+ chunks at once crashed Ollama's
# /api/embed with a connection reset after ~30s. Batching client-side avoids
# that -- 64 matches what a production batch size would reasonably look like.


def run_ingestion(files: List[Path]):
    manifest = load_manifest()
    chunks = load_chunks()
    vectorstore = None
    if os.path.exists(FAISS_DIR):
        vectorstore = FAISS.load_local(FAISS_DIR, get_embeddings(), allow_dangerous_deserialization=True)

    to_process = [(p, file_hash(p)) for p in files]
    to_process = [(p, h) for p, h in to_process if manifest.get(str(p), {}).get("hash") != h]

    print(f"[ingest] {len(files)} file(s) scanned, {len(to_process)} new/changed")
    if not to_process:
        print("[ingest] nothing to do -- index already up to date.")
        return

    embeddings = get_embeddings()
    for path, h in to_process:
        path_str = str(path)
        old_entry = manifest.get(path_str)
        if old_entry:
            if vectorstore is not None and old_entry["chunk_ids"]:
                vectorstore.delete(ids=old_entry["chunk_ids"])
            chunks = [c for c in chunks if c.metadata.get("source_path") != path_str]
            print(f"[ingest] removed {len(old_entry['chunk_ids'])} stale chunk(s) for changed file: {path.name}")

        file_chunks = build_chunks([path])
        chunk_ids = [str(uuid.uuid4()) for _ in file_chunks]

        for i in range(0, len(file_chunks), EMBED_BATCH_SIZE):
            batch = file_chunks[i:i + EMBED_BATCH_SIZE]
            batch_ids = chunk_ids[i:i + EMBED_BATCH_SIZE]
            if vectorstore is None:
                vectorstore = FAISS.from_documents(batch, embeddings, ids=batch_ids)
            else:
                vectorstore.add_documents(batch, ids=batch_ids)

        chunks.extend(file_chunks)
        manifest[path_str] = {"hash": h, "chunk_ids": chunk_ids}
        print(f"[ingest] embedded {len(file_chunks)} chunk(s) from {path.name}")

    if vectorstore is not None:
        vectorstore.save_local(FAISS_DIR)
    save_chunks(chunks)
    save_manifest(manifest)
    print(f"[ingest] done. {len(chunks)} total chunk(s) indexed, {len(to_process)} file(s) updated this run.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--scan-dir", default=os.environ.get("LOCAL_SCAN_DIR", "~/Downloads"))
    parser.add_argument("--extensions", default=None,
                         help="Comma-separated list to restrict ingestion to (e.g. 'pdf,docx'). Default: all supported types.")
    args = parser.parse_args()

    scan_dir = Path(args.scan_dir).expanduser()
    all_files = collect_files(scan_dir)

    if args.extensions:
        allowed = {f".{e.strip().lower().lstrip('.')}" for e in args.extensions.split(",")}
        all_files = [f for f in all_files if f.suffix.lower() in allowed]

    print(f"[ingest] found {len(all_files)} file(s) in {scan_dir}")

    if not all_files:
        print("[ingest] nothing to ingest -- add files to the scan folder and re-run.")
    else:
        run_ingestion(all_files)
